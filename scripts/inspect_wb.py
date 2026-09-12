import openpyxl

wb = openpyxl.load_workbook('./public/templates/MPSS BLANK REPORT.xlsx', data_only=False)
print("Sheets in workbook:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n================ Sheet: {sheetname} | max_row: {ws.max_row} | max_col: {ws.max_column} ================")
    merged = [str(m) for m in list(ws.merged_cells.ranges)[:10]]
    if merged:
        print(f"  Merged ranges (first 10): {', '.join(merged)}")

    for r in range(1, min(ws.max_row + 1, 40)):
        row_vals = []
        for c in range(1, min(ws.max_column + 1, 15)):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                row_vals.append(f"{ws.cell(row=r, column=c).coordinate}:{str(val)[:30]}")
        if row_vals:
            print(f"  Row {r}: {' | '.join(row_vals)}")
